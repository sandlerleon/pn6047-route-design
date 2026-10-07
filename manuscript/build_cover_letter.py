# -*- coding: utf-8 -*-
"""Cover letter for the submission to the Journal of Pharmaceutical Innovation (original submission); numbers read from results/results.json.

    python build_cover_letter.py     ->  out/Cover_Letter_JPI_v1.docx
"""
import json
import os
import sys

from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx_helpers as H  # noqa: E402

R = json.load(open(os.path.join(HERE, "..", "results", "results.json"), encoding="utf-8"))
ZEN = json.load(open(os.path.join("C:" + os.sep, "YouTube", "_pn6047_zenodo_state.json")))
SW, PP = ZEN["software"]["doi"], ZEN["publication"]["doi"]
REPO = "https://github.com/sandlerleon/pn6047-route-design"
D = {k: v["default"] for k, v in R["routes"].items()}
TITLE = "A Literature-Based Modular Route-Design Strategy for Scalable Manufacture of PN6047"
doc = H.new_document(size=11, line=1.15)


def para(text, bold=False, after=6):
    p = doc.add_paragraph()
    if bold:
        p.add_run(text).bold = True
    else:
        H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(3)


for line in ("Leon Sandler", "Independent researcher, Northbrook, Illinois, USA", "sandler.leon@gmail.com", "ORCID: https://orcid.org/0009-0007-4584-808X"):
    para(line, after=0)
para("")
para("6 October 2026")
para("The Editor\nJournal of Pharmaceutical Innovation", after=10)
para("Submission of a Research Article: “%s”" % TITLE, bold=True, after=10)
para("Dear Editor,")
para("I submit the enclosed paper for consideration in the Journal of Pharmaceutical Innovation as a Research Article. It asks a question that arises before any "
     "process chemistry is run on a new drug substance: which operations of the only publicly exemplified route would limit manufacture, and how strong is the public "
     "evidence for the alternatives that would replace them? The case is PN6047, a clinical-stage δ-opioid receptor agonist whose discovery-scale route is public only as "
     "milligram-to-gram patent examples.")
para("What the paper does", bold=True, after=3)
bullet("It inventories the operations exemplified in WO 2016/099393 and WO 2022/013153 and grades each transformation, and its conditions, in four evidence classes "
       "(direct, close precedent, general precedent, hypothesis), so that what is established, inferred and proposed is never mixed.")
bullet("It flags documented regulatory and hazard attributes of the exemplified route (%d flag instances, including an ICH Q3C Class 1 solvent, palladium, a "
       "benzotriazole-derived reagent and preparative HPLC) and screens two process-oriented variants: they carry %d flags each, but their validation burden is higher "
       "(%d and %d against %d on the linear scale), because their conditions are untested." % (D["R0"]["flag_instances"], D["R1"]["flag_instances"],
       D["R1"]["validation_burden"]["linear"], D["R2"]["validation_burden"]["linear"], D["R0"]["validation_burden"]["linear"]))
bullet("It tests how robust the ranking of a twelve-item risk register is to the author's own uncertainty; two risks, impurity purge when crystallisation replaces "
       "chromatography and the undisclosed upstream route to the vinyl bromide building block, dominate.")
bullet("It ends with a stage-gated experimental programme, and states what is not claimed: no synthesis was performed, and no yield, impurity level, cost or timeline is predicted.")
para("The paper also corrects a mapping on which an earlier outline of the idea rested: pramipexole is a useful benchmark for manufacture of a basic thiazole-containing "
     "API, but its benzothiazole core is not a fragment of PN6047.", after=6)
para("Why the journal", bold=True, after=3)
para("The study concerns route selection, process risk, impurity control, starting-material justification (ICH Q11) and the evidence needed before scale-up, which I believe fall "
     "within the journal's coverage of manufacturing and applied pharmaceutical science, and it is offered as a method that can be applied to other compounds whose "
     "discovery routes are public. It follows the Research Article format and the abstract, summary and teaser requirements.", after=6)
para("Data, code and disclosures", bold=True, after=3)
para("All inputs, model code, tests and figure scripts are public at %s and archived at https://doi.org/%s; the manuscript is available as a preprint at "
     "https://doi.org/%s. A large language model (Claude, Anthropic) was used to assist with retrieval, code, figures and drafting; I reviewed and verified the content, "
     "resolved every reference through Crossref, and take full responsibility, as stated in Section 2.9. I am the sole author, have no competing interests, received no "
     "funding and have no affiliation with the applicant of the patents analysed. An earlier, shorter and unpublished outline of the route concept was submitted to an "
     "open-innovation challenge, as disclosed in the Declarations. The manuscript has not been published and is not under consideration by any other journal." % (REPO, SW, PP))
para("Thank you for considering the paper.")
para("Yours sincerely,", after=18)
para("Leon Sandler")
out = os.path.join(HERE, "out", "Cover_Letter_JPI_v1.docx")
doc.save(out)
print("saved", out)
